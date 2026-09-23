import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx

from src.docx_extract import extract_docx


def test_extract_docx_returns_text():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "synthetic.docx")
        document = docx.Document()
        document.add_paragraph("Jane Doe")
        document.add_paragraph("Python developer with a LangChain RAG project.")
        document.save(path)

        result = extract_docx(path)

        assert result.ok is True
        assert "Jane Doe" in result.text
        assert "LangChain" in result.text


def test_extract_docx_handles_corrupt_file_without_crashing():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "corrupt.docx")
        with open(path, "wb") as f:
            f.write(b"not a real docx file")

        result = extract_docx(path)

        assert result.ok is False
        assert result.error is not None
