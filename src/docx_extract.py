"""Deterministic DOCX extraction (bonus, optional per assignment).

Mirrors the shape of pdf_extract.PdfExtractionResult so pipeline.py can treat
PDF and DOCX inputs uniformly. Hyperlink-annotation extraction (the PDF-only
trick used for GitHub links) is not attempted here — DOCX candidates rely on
visible-text GitHub URLs only, which extract_fields.py already handles.
"""
from __future__ import annotations

from dataclasses import dataclass, field

import docx


@dataclass
class DocxExtractionResult:
    ok: bool
    text: str = ""
    hyperlinks: list[str] = field(default_factory=list)
    error: str | None = None


def extract_docx(path: str) -> DocxExtractionResult:
    try:
        document = docx.Document(path)
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(c.text for c in row.cells)
        text = "\n".join(parts)
    except Exception as e:  # noqa: BLE001 - one bad file must not crash the batch
        return DocxExtractionResult(ok=False, error=f"docx_extraction_failed: {e}")

    if not text.strip():
        return DocxExtractionResult(ok=False, error="empty_text_extraction")

    return DocxExtractionResult(ok=True, text=text, hyperlinks=[])
